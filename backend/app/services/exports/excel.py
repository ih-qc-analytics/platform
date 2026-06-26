from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any, Callable, Mapping

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


PostProcessHook = Callable[[Worksheet, list["ExcelColumn"], list[Mapping[str, Any]]], None]

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E2F3"),
    right=Side(style="thin", color="D9E2F3"),
    top=Side(style="thin", color="D9E2F3"),
    bottom=Side(style="thin", color="D9E2F3"),
)


@dataclass(frozen=True)
class ExcelColumn:
    key: str
    header: str


@dataclass(frozen=True)
class ExcelWorksheetSpec:
    name: str
    columns: list[ExcelColumn]
    rows: list[Mapping[str, Any]]
    post_process: PostProcessHook | None = None
    note: str | None = None  # rendered as a styled row above the column headers


def _safe_sheet_name(name: str, existing: set[str]) -> str:
    base_name = (name or "Sheet").strip()[:31] or "Sheet"
    candidate = base_name
    suffix = 1
    while candidate in existing:
        suffix_text = f" {suffix}"
        candidate = f"{base_name[: 31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    return candidate


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(item) for item in value)
    return value


def _style_header_row(worksheet: Worksheet, header_count: int, header_row: int = 1) -> None:
    for column_index in range(1, header_count + 1):
        cell = worksheet.cell(row=header_row, column=column_index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def _apply_body_styles(worksheet: Worksheet, data_start_row: int = 2) -> None:
    for row in worksheet.iter_rows(min_row=data_start_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER


def _autosize_columns(worksheet: Worksheet) -> None:
    for column_cells in worksheet.columns:
        values = [len(str(cell.value or "")) for cell in column_cells]
        if not values:
            continue
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
            max(max(values) + 2, 12),
            60,
        )


def build_excel_workbook(worksheets: list[ExcelWorksheetSpec]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    existing_names: set[str] = set()

    for spec in worksheets:
        sheet_name = _safe_sheet_name(spec.name, existing_names)
        existing_names.add(sheet_name)
        worksheet = workbook.create_sheet(title=sheet_name)

        header_row = 1
        if spec.note:
            note_cell = worksheet.cell(row=1, column=1, value=spec.note)
            note_cell.font = Font(italic=True, color="6B7280", size=9)
            note_cell.alignment = Alignment(vertical="center")
            if len(spec.columns) > 1:
                worksheet.merge_cells(
                    start_row=1,
                    start_column=1,
                    end_row=1,
                    end_column=len(spec.columns),
                )
            worksheet.row_dimensions[1].height = 16
            header_row = 2

        worksheet.append([column.header for column in spec.columns])
        for row in spec.rows:
            worksheet.append([_cell_value(row.get(column.key)) for column in spec.columns])

        freeze_col = get_column_letter(1)
        worksheet.freeze_panes = f"{freeze_col}{header_row + 1}"
        _style_header_row(worksheet, len(spec.columns), header_row=header_row)
        _apply_body_styles(worksheet, data_start_row=header_row + 1)
        _autosize_columns(worksheet)

        if spec.post_process is not None:
            spec.post_process(worksheet, spec.columns, spec.rows)

    return workbook


def workbook_to_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def generate_excel_response(
    filename: str, worksheets: list[ExcelWorksheetSpec]
) -> StreamingResponse:
    workbook = build_excel_workbook(worksheets)
    workbook_bytes = workbook_to_bytes(workbook)
    return StreamingResponse(
        iter([workbook_bytes]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
    )


def add_total_sales_charts(
    worksheet: Worksheet,
    columns: list[ExcelColumn],
    rows: list[Mapping[str, Any]],
) -> None:
    if len(rows) <= 1:
        return

    col_keys = [c.key for c in columns]
    has_comparison = "comp_trend_revenue" in col_keys
    data_rows = len(rows)

    # Detect whether there's a note row above the column headers
    header_row = 1 if worksheet.cell(row=1, column=1).value == columns[0].header else 2
    data_start = header_row + 1
    data_end = header_row + data_rows

    trend_col = col_keys.index("trend_revenue") + 1
    dim_col = col_keys.index("dimension") + 1
    geo_col = col_keys.index("geo_revenue") + 1

    trend_chart = LineChart()
    trend_chart.title = "Revenue Trend"
    trend_chart.y_axis.title = "Revenue"
    trend_chart.x_axis.title = "Month"
    trend_data = Reference(worksheet, min_col=trend_col, min_row=header_row, max_row=data_end)
    trend_categories = Reference(worksheet, min_col=1, min_row=data_start, max_row=data_end)
    trend_chart.add_data(trend_data, titles_from_data=True)
    trend_chart.set_categories(trend_categories)
    if has_comparison:
        comp_trend_col = col_keys.index("comp_trend_revenue") + 1
        comp_trend_data = Reference(
            worksheet, min_col=comp_trend_col, min_row=header_row, max_row=data_end
        )
        trend_chart.add_data(comp_trend_data, titles_from_data=True)
    worksheet.add_chart(trend_chart, "G2")

    geo_chart = BarChart()
    geo_chart.title = "Revenue by Country"
    geo_chart.y_axis.title = "Revenue"
    geo_chart.x_axis.title = "Country"
    geo_data = Reference(worksheet, min_col=geo_col, min_row=header_row, max_row=data_end)
    geo_categories = Reference(worksheet, min_col=dim_col, min_row=data_start, max_row=data_end)
    geo_chart.add_data(geo_data, titles_from_data=True)
    geo_chart.set_categories(geo_categories)
    if has_comparison:
        comp_geo_col = col_keys.index("comp_geo_revenue") + 1
        comp_geo_data = Reference(
            worksheet, min_col=comp_geo_col, min_row=header_row, max_row=data_end
        )
        geo_chart.add_data(comp_geo_data, titles_from_data=True)
    worksheet.add_chart(geo_chart, "G20")
